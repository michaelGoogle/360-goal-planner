"""WP0: export parity fixtures from the calculation workbook.

Recalculates the workbook for every persona (and a few scenarios) with LibreOffice and writes
one JSON per case with inputs and the expected output of every component.

    python export_fixtures.py [--workbook "Calculations V0-24.xlsx"] [--out ../../tests/fixtures/model_v0_24]
                              [--work /tmp/wp0] [--only P01,P03]

Needs: openpyxl, LibreOffice (soffice) on PATH. Does not modify the workbook.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import math
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
MODEL = "V0-24"
SNAPSHOT_SGD = 0.78162734281267
SNAPSHOT_VND = 3.849039994550423e-05

GROUPS = {
    "fx": "FX_", "cpf": "CPF_", "peopleLikeYou": "PLU_", "needProfiler": "NP_",
    "needCalculator": "NC_", "plan": "PLAN_", "happiU": "HU_", "scenarioVisualizer": "SV_",
}


def _clean(v):
    if isinstance(v, float):
        if math.isnan(v) or math.isinf(v):
            return None
        return v
    if isinstance(v, (dt.date, dt.datetime)):
        return v.isoformat()
    return v


def _names(wb) -> dict:
    out = {}
    for n, dn in wb.defined_names.items():
        try:
            dests = list(dn.destinations)
        except Exception:
            continue
        if len(dests) != 1:
            continue
        sh, ref = dests[0]
        if sh not in wb.sheetnames:
            continue
        rng = wb[sh][ref.replace("$", "")]
        if isinstance(rng, tuple):
            flat = [c for row in rng for c in (row if isinstance(row, tuple) else (row,))]
            out[n] = [_clean(c.value) for c in flat]
        else:
            out[n] = _clean(rng.value)
    return out


def _label_block(ws, labels_prefix: tuple[str, ...]) -> dict:
    out = {}
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, str) and a.startswith(labels_prefix):
            out[a] = _clean(ws.cell(r, 2).value)
    return out


def _labels_all(ws) -> dict:
    out = {}
    for r in range(1, ws.max_row + 1):
        a, b, k = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value
        if isinstance(a, str) and isinstance(k, str) and k[:1] in "ABCDE" and "·" in k:
            out[a] = _clean(b)
    return out


def extract(path: Path, case: str, edits: dict) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    names = _names(wb)
    doc = {"case": case, "model": MODEL, "fxSnapshot": "2026-09-25", "edits": {f"{k[0]}!{k[1]}": v for k, v in edits.items()}}
    ws = wb["Personas"]
    pid = wb["Selected persona"]["B2"].value
    heads = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
    for r in range(3, ws.max_row + 1):
        if ws.cell(r, 1).value == pid:
            doc["inputs"] = {str(h): _clean(ws.cell(r, c).value) for c, h in enumerate(heads, 1) if h}
            break
    for g, pre in GROUPS.items():
        scal = {k: v for k, v in names.items() if k.startswith(pre) and not isinstance(v, list)}
        arr = {k: v for k, v in names.items() if k.startswith(pre) and isinstance(v, list)}
        doc[g] = {"values": dict(sorted(scal.items()))}
        if arr:
            doc[g]["years"] = dict(sorted(arr.items()))
    np_ws = wb["Need Profiler"]
    doc["needProfiler"]["raw"] = _label_block(np_ws, ("raw ",))
    doc["needProfiler"]["scaled"] = _label_block(np_ws, ("scaled ",))
    doc["needProfiler"]["optionWeights"] = _label_block(np_ws, ("ow ", "usd "))
    doc["budget"] = {"values": _labels_all(wb["Budget Calculator"])}
    doc["plan"]["rows"] = _labels_all(wb["Plan Calculator"])
    doc["needCalculator"]["rows"] = _labels_all(wb["Need Calculator"])
    errs = []
    for sh in wb.worksheets:
        for row in sh.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value[:4] in ("#REF", "#VAL", "#NAM", "#DIV", "#N/A", "#NUM", "Err:") \
                        and not (sh.title == "FX" and c.column == 4):
                    errs.append(f"{sh.title}!{c.coordinate}={c.value}")
    doc["formulaErrors"] = errs
    return doc


def parameters(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Assumptions"]
    out = {"model": MODEL, "parameters": {}}
    for n, dn in wb.defined_names.items():
        try:
            dests = list(dn.destinations)
        except Exception:
            continue
        if len(dests) != 1 or dests[0][0] != "Assumptions":
            continue
        ref = dests[0][1].replace("$", "")
        cell = ws[ref]
        if isinstance(cell, tuple):
            continue
        r = cell.row
        meta = {}
        if cell.column == 2:
            meta = {k: ws.cell(r, c).value for k, c in (("unit", 3), ("currency", 4), ("status", 5),
                                                        ("codeSource", 6), ("changedIn", 7), ("openProposal", 8))}
            meta["label"] = ws.cell(r, 1).value
        out["parameters"][n] = {"value": _clean(cell.value), "cell": ref, **{k: v for k, v in meta.items() if v not in (None, "")}}
    out["parameters"] = dict(sorted(out["parameters"].items()))
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook", default=f"Calculations {MODEL}.xlsx")
    ap.add_argument("--out", default=str(HERE.parents[1] / "tests" / "fixtures" / "model_v0_24"))
    ap.add_argument("--work", default="/tmp/wp0")
    ap.add_argument("--only", default="")
    ap.add_argument("--phase", default="all", choices=["prepare", "calc", "extract", "all"])
    a = ap.parse_args()
    src = (HERE / a.workbook).resolve()
    work = Path(a.work)
    (work / "in").mkdir(parents=True, exist_ok=True)
    (work / "out").mkdir(parents=True, exist_ok=True)
    out = Path(a.out)
    (out / "personas").mkdir(parents=True, exist_ok=True)
    (out / "scenarios").mkdir(parents=True, exist_ok=True)

    base = openpyxl.load_workbook(src)
    ids = [base["Personas"].cell(r, 1).value for r in range(3, 53)]
    row_of = {pid: 3 + i for i, pid in enumerate(ids)}
    only = [x for x in a.only.split(",") if x]
    cases: dict[str, dict] = {}
    for pid in ids:
        if not only or pid in only:
            cases[f"personas/{pid}"] = {("Selected persona", "B2"): pid}

    def cell_of(name: str) -> tuple[str, str]:
        (sh, ref), = list(base.defined_names[name].destinations)
        return sh, ref.replace("$", "")

    if not only:
        r3 = row_of["P03"]
        k = SNAPSHOT_SGD / SNAPSHOT_VND
        income = base["Personas"][f"J{r3}"].value
        cases["scenarios/P03_VND_Vietnam"] = {("Selected persona", "B2"): "P03", ("Personas", f"F{r3}"): "Vietnam",
                                              ("Personas", f"K{r3}"): "VND", ("Personas", f"J{r3}"): income * k}
        e = {("Selected persona", "B2"): "P03"}
        for t in ("N_PAC", "N_LTC"):
            sh, ref = cell_of(f"NC_on_{t}")
            e[(sh, "E" + "".join(ch for ch in ref if ch.isdigit()))] = True
        cases["scenarios/P03_PAC_LTC_on"] = e
        e = {("Selected persona", "B2"): "P03"}
        for kk in ("dea", "cri", "tpd", "pac", "hos", "ltc", "wed", "bab", "lon"):
            e[cell_of(f"SV_{kk}V")] = True
        cases["scenarios/P03_all_stress_events"] = e

    def fname(case: str) -> str:
        return case.replace("/", "__")

    if a.phase in ("prepare", "all"):
        for case, edits in cases.items():
            wb = openpyxl.load_workbook(src)
            for (sh, ref), v in edits.items():
                wb[sh][ref] = v
            wb.save(work / "in" / f"{fname(case)}.xlsx")
            print("prepared", case, flush=True)
    if a.phase in ("calc", "all"):
        files = sorted(str(p) for p in (work / "in").glob("*.xlsx"))
        for i in range(0, len(files), 8):
            chunk = files[i:i + 8]
            subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
                            "--outdir", str(work / "out"), *chunk], capture_output=True, timeout=1800)
            print("calculated", i + len(chunk), "of", len(files), flush=True)
    if a.phase in ("extract", "all"):
        (out / "parameters.json").write_text(json.dumps(parameters(work / "out" / f"{fname(next(iter(cases)))}.xlsx"),
                                                        indent=1, ensure_ascii=False, default=str), encoding="utf-8")
        index = []
        for case, edits in cases.items():
            f = work / "out" / f"{fname(case)}.xlsx"
            if not f.exists():
                print("MISSING", case, flush=True)
                continue
            doc = extract(f, case.split("/")[1], edits)
            (out / f"{case}.json").write_text(json.dumps(doc, indent=1, ensure_ascii=False, default=str), encoding="utf-8")
            index.append({"case": case, "formulaErrors": len(doc["formulaErrors"])})
            print("extracted", case, len(doc["formulaErrors"]), flush=True)
        (out / "index.json").write_text(json.dumps({"model": MODEL, "workbook": src.name, "cases": index}, indent=1),
                                        encoding="utf-8")
    return 0


if __name__ == "__main__":
    sys.exit(main())
